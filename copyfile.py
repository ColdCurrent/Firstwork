# -*- coding: utf-8 -*-
"""
Created on Thu Apr  9 22:52:14 2020

@author: Huang
"""


import os
import shutil as sh
from docx import Document


pathdr = os.getcwd()   #資料夾的位置 "C:\Users\阿翔\Desktop\Word Automation"
ls = os.listdir(pathdr)   #資料夾所在位置的資料夾名稱的array ["001黃翔翔1"]
st = ls[0]                   #"001黃翔翔1"
               
pathfile = pathdr +"\\"+st   #"C:\Users\阿翔\Desktop\Word Automation\001黃翔翔1"
name_of_file = os.listdir(pathfile)[0]  #目標檔案名"001黃翔翔1-清償提存書.docx"


for i in range(2,5):

    sttemp = name_of_file.replace("1",str(i))   #"00i黃翔翔i-清償提存書.docx"
    pathtemp = pathdr+"\\"+st.replace("1",str(i)) #要創立的資料夾名稱"C:\Users\阿翔\Desktop\Word Automation\00i黃翔翔i"

    sh.copytree(pathfile,pathtemp)
    os.rename(os.path.join(pathtemp,name_of_file) , os.path.join(pathtemp,sttemp))
    

    path = os.path.join(pathtemp,sttemp)
    
    
    document = Document(path)
    tables = document.tables     #表格做成的ARRAY
    table = tables[0] 
    st2 = table.cell(2,1).text
    a=st2.split( )

    a1 = a[0].replace("1",str(i))
    a2 = a[1][:-3]+str(int(a[1][-3:])+i-1)

    table.cell(2,1).text = a1 + "\n" + a2
    table.cell(2,3).text = table.cell(2,3).text.replace("1",str(i))
    document.save(path)